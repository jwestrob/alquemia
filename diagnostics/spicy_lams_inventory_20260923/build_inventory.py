"""Read permitted metadata only; organize existing files without molecular analysis."""
import argparse,collections,csv,hashlib,json,re,subprocess
from pathlib import Path
import openpyxl,gemmi

ROOT=Path(__file__).resolve().parents[2]
JACOB=Path('/groups/banfield/projects/environmental/sr/srvp2020/Jacob')
PROJECT=JACOB/'spicy_lams'
BINDING=JACOB/'lanthanide_binding'
SENTINEL=ROOT/'spicy_lams_matched_homolog'
LEFT=('ID-4','ID-6','nodeacc1','nodeacc2','db','Sequence..SP.Removed.','source_original','phylum_original','class_original','order_original','family_original','genus_original','species_original','EFhands','mxaF','xoxF','exaF','mGDH','metabolism')
QC=('Total_Metal_Replicate_1','Total_Metal_Replicate_2','Total_Metal_Replicate_3','Total_Metal_Replicate_Average','Total_Metal_Replicate_AverageRounded','Total_Metal_Replicate_SD','Total_Metal_Replicate_RSD')

def pin(p):
 p=Path(p);h=hashlib.sha256()
 with p.open('rb')as f:
  for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
 return {'path':str(p.resolve()),'sha256':h.hexdigest(),'bytes':p.stat().st_size}
def seqhash(s):return hashlib.sha256(s.encode()).hexdigest()
def write(out,name,data):
 p=out/name
 with p.open('x')as f:json.dump(data,f,indent=2,sort_keys=True);f.write('\n')
 return pin(p)
def table(out,name,rows):
 p=out/name
 with p.open('x',newline='')as f:
  if rows:
   w=csv.DictWriter(f,fieldnames=list(rows[0]),delimiter='\t');w.writeheader()
   for row in rows:w.writerow({k:json.dumps(v,sort_keys=True)if isinstance(v,(dict,list))else v for k,v in row.items()})
 return pin(p)
def filelist(root):
 r=subprocess.run(['rg','--files','--hidden',str(root)],text=True,capture_output=True)
 if r.returncode not in (0,1):raise RuntimeError(r.stderr)
 return sorted(Path(p)for p in r.stdout.splitlines())
def structure(p):
 s=gemmi.read_structure(str(p));chains=[];hetero=collections.Counter();models=len(s)
 for c in s[0]:
  residues=[r for r in c if gemmi.find_tabulated_residue(r.name).is_amino_acid()]
  seq=''.join(gemmi.find_tabulated_residue(r.name).one_letter_code for r in residues)
  if seq:chains.append({'chain':c.name,'sequence':seq,'sequence_sha256':seqhash(seq),'residues':len(residues)})
  for r in c:
   if not gemmi.find_tabulated_residue(r.name).is_amino_acid():hetero[r.name]+=1
 return {'file':pin(p),'models':models,'chains':chains,'first_model_hetero_residues':dict(hetero)}
def main(output):
 out=Path(output).resolve();out.mkdir(parents=True,exist_ok=False)
 workbook=PROJECT/'41589_2026_2176_MOESM3_ESM.xlsx';wb=openpyxl.load_workbook(workbook,read_only=True,data_only=True);ws=wb['Data']
 # Never iterate a combined A:BS block: that would expose forbidden T:BL.
 left=list(ws.iter_rows(min_col=1,max_col=19,values_only=True));qc=list(ws.iter_rows(min_col=65,max_col=71,values_only=True))
 if tuple(left[0])!=LEFT or tuple(qc[0])!=QC or len(left)!=len(qc):raise ValueError('allowed metadata schema changed')
 selection=json.loads((SENTINEL/'selection_manifest.json').read_text());reserved={m['accession']:(p['pair_id'] if 'pair_id'in p else f'p{i+1:02}',j+1,m)for i,p in enumerate(selection['pairs'])for j,m in enumerate(p['members'])}
 download=list(csv.DictReader((PROJECT/'metadata/lanm_download_summary.tsv').open(),delimiter='\t'));dl={r['id4']:r for r in download}
 records=[]
 for idx,(a,q)in enumerate(zip(left[1:],qc[1:]),2):
  if not any(x is not None for x in a):continue
  r=dict(zip(LEFT,a));r.update(dict(zip(QC,q)));seq=str(r['Sequence..SP.Removed.']).replace(' ','').strip();acc=r['nodeacc2'];d=dl.get(r['ID-4'],{})
  if d and d['protein_acc']!=acc:raise ValueError('metadata accession join differs')
  row={'workbook_row':idx,'id4':r['ID-4'],'id6':r['ID-6'],'ortholog_id':'o-'+str(r['ID-6']),'accession':acc,'database':r['db'],
       'mature_sequence':seq,'sequence_length':len(seq),'sequence_sha256':seqhash(seq),'source_organism':r['source_original'],
       'family':r['family_original'],'genus':r['genus_original'],'ef_hand_motifs':r['EFhands'],'reserved_panel':acc in reserved,
       'reserved_pair':reserved[acc][0]if acc in reserved else None,'reserved_member':reserved[acc][1]if acc in reserved else None,
       'assay_outcomes':'not_read','construct_scope':'mature LanM sequence; experimental SpyTag/GSG/immobilization not modeled',
       'assembly_acc':d.get('assembly_acc')or None}
  if acc in reserved and(seq!=reserved[acc][2]['mature_sequence']or idx!=reserved[acc][2]['source_row']):raise ValueError('reserved sequence/row changed')
  for kind in ('genome','proteome','annotation'):
   stored=d.get(kind)or None;located=PROJECT/Path(stored).parent.name/Path(stored).name if stored else None
   row[kind+'_original_path']=stored;row[kind+'_original_exists']=Path(stored).is_file()if stored else False
   row[kind+'_located_path']=str(located)if located and located.is_file()else None
  records.append(row)
 counts=collections.Counter(r['sequence_sha256']for r in records)
 for r in records:r['same_mature_sequence_rows']=counts[r['sequence_sha256']]
 artifacts={'orthologs':table(out,'ORTHOLOGS.tsv',records)}
 with (out/'MATURE_SEQUENCES.faa').open('x')as f:
  for r in records:f.write(f">{r['ortholog_id']}|{r['accession']}|id4={r['id4']}|reserved={str(r['reserved_panel']).lower()}\n{r['mature_sequence']}\n")
 artifacts['sequences']=pin(out/'MATURE_SEQUENCES.faa')
 prediction=json.loads((SENTINEL/'af2_prediction_manifest.json').read_text());folds=[];structures=[]
 for p in prediction['predictions']:
  acc=next(a for a in reserved if a in p['name']);q=next(r for r in records if r['accession']==acc)
  base=SENTINEL/'folding_runs/af2_template_v1/predict_out'/p['name'];models=sorted(base.glob('*_unrelaxed_rank_*.pdb'))
  model_rows=[]
  for model in models:
   meta=structure(model);meta.update(dataset='reserved_AF2',accession=acc,ortholog_id=q['ortholog_id'])
   if len(meta['chains'])!=1 or meta['chains'][0]['sequence']!=q['mature_sequence']:raise ValueError('full mature-chain model differs')
   meta['mature_sequence_exact']=True;structures.append(meta);model_rows.append(meta['file'])
  folds.append({'accession':acc,'ortholog_id':q['ortholog_id'],'reserved_pair':q['reserved_pair'],'reserved_member':q['reserved_member'],
   'name':p['name'],'fasta':p['fasta'],'locked_msa':p['locked_a3m'],'models':model_rows,'model_count':len(models),
   'full_mature_chain_exact':bool(models),'completion_marker_present':bool(list(base.glob('*.done.txt'))),
   'confidence_paths':[str(f.resolve())for f in sorted(base.glob('*scores_rank*.json'))],
   'scope':'protein-only AF2 template-conditioned models; no transferred metals/waters/core preparation inferred'})
 artifacts['folds']=table(out,'RESERVED_FOLDS.tsv',folds)
 other_roots={'Sharur_discovery_candidates':BINDING/'candidate_bundle/spici_lams_other',
              'FSRD_discovery_LanM_calibration':BINDING/'candidate_bundle/lanm_calibration',
              'LanM_deposited_sources':BINDING/'lanm_benchmark/structures',
              'LanM_four_La_prediction_trial':BINDING/'folding_runs/lanm_4La_multimer_test/predict_out'}
 for dataset,base in other_roots.items():
  for p in filelist(base):
   if p.suffix.lower()not in('.pdb','.cif'):continue
   meta=structure(p);matches=sorted({r['accession']for r in records for c in meta['chains']if r['sequence_sha256']==c['sequence_sha256']})
   meta.update(dataset=dataset,exact_workbook_mature_sequence_matches=matches);structures.append(meta)
 meta=structure(SENTINEL/'templates/8FNS.pdb');meta.update(dataset='reserved_template_only');structures.append(meta)
 for p in sorted((BINDING/'fep/prep').glob('*.pdb')):
  meta=structure(p);meta.update(dataset='historical_full_protein_FEP_preparation');structures.append(meta)
 artifacts['structures']=write(out,'STRUCTURES.json',structures)
 project_files=[]
 for p in filelist(PROJECT):
  project_files.append({'path':str(p),'bytes':p.stat().st_size,'category':str(p.relative_to(PROJECT)).split('/')[0],'is_symlink':p.is_symlink(),'resolved_path':str(p.resolve())})
 artifacts['project_files']=table(out,'PROJECT_FILES.tsv',project_files)
 legacy_roots={'LanM_benchmark':BINDING/'lanm_benchmark','LanM_FEP':BINDING/'fep','Hans_Protenix':ROOT/'hans_lanm_protenix_benchmark',
               'LanM_MACE_series':ROOT/'workspaces/lanm_series_followup_20260923','reserved_AF2':SENTINEL}
 reps=[]
 for dataset,base in legacy_roots.items():
  files=filelist(base)
  kept=[p for p in files if p.suffix.lower()in('.pdb','.cif','.dcd','.xtc','.prmtop','.parm7','.top','.xml','.faa','.fasta')]
  for p in kept:reps.append({'dataset':dataset,'path':str(p),'bytes':p.stat().st_size,'type':p.suffix,'contents_evaluated':False})
 artifacts['representation_files']=table(out,'REPRESENTATION_FILES.tsv',reps)
 summary={'workbook':pin(workbook),'allowed_ranges':['Data!A:S','Data!BM:BS'],'forbidden_ranges_not_read':['Data!T:BL','Data!BT:CW','Data!CX:CZ'],
  'selection':pin(SENTINEL/'selection_manifest.json'),'prediction_manifest':pin(SENTINEL/'af2_prediction_manifest.json'),
  'workbook_ortholog_rows':len(records),'unique_accessions':len({r['accession']for r in records}),'unique_mature_sequences':len(counts),
  'duplicate_sequence_groups':[{'sha256':h,'members':[r['ortholog_id']for r in records if r['sequence_sha256']==h]}for h,n in counts.items()if n>1],
  'mature_length_range':[min(r['sequence_length']for r in records),max(r['sequence_length']for r in records)],
  'ef_hand_annotation_counts':dict(collections.Counter(str(r['ef_hand_motifs'])for r in records)),
  'reserved_proteins':len(folds),'reserved_full_chain_model_count':sum(r['model_count']for r in folds),'reserved_all_completed':all(r['completion_marker_present']for r in folds),
  'reserved_model_hetero_inventory':dict(collections.Counter(tuple(sorted(m['first_model_hetero_residues']))for m in structures if m['dataset']=='reserved_AF2')),
  'assembly_mapped_rows':sum(bool(r['assembly_acc'])for r in records),'unique_associated_assemblies':len({r['assembly_acc']for r in records if r['assembly_acc']}),
  'located_genome_rows':sum(bool(r['genome_located_path'])for r in records),'located_proteome_rows':sum(bool(r['proteome_located_path'])for r in records),
  'stale_original_proteome_paths':sum(bool(r['proteome_original_path'])and not r['proteome_original_exists']for r in records),
  'project_category_counts':dict(collections.Counter(r['category']for r in project_files)),
  'representation_counts':dict(collections.Counter(r['dataset']for r in reps)),'artifacts':artifacts,
  'new_molecular_calls':0,'outcomes_read':False,'source_files_modified':False,'limitations':'filename inventory is bounded to declared roots; absence elsewhere is not proven; no new site or MD analysis'}
 summary['reserved_model_hetero_inventory']={str(k):v for k,v in summary['reserved_model_hetero_inventory'].items()}
 write(out,'SUMMARY.json',summary);print(json.dumps({k:v for k,v in summary.items()if k not in ['artifacts','duplicate_sequence_groups']},indent=2))
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--output',required=True);main(p.parse_args().output)
